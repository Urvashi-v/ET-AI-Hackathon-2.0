"""Word document parser.

Incident reports and investigation narratives arrive as ``.docx`` more often
than anything else. Word's own style names give a reliable heading hierarchy, so
``section_path`` comes out better here than from a PDF.

Tables are emitted as ``table_row`` blocks with their header repeated per row.
Flattening a table into prose destroys the row-column association, and a model
reading a flattened thickness table will confidently read the wrong cell.
"""

from __future__ import annotations

import io

from services.ingest.parsers.base import ParsedDocument, TextBlock


class DocxParser:
    name = "python-docx"

    def can_parse(self, extension: str) -> bool:
        return extension.lower() == ".docx"

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        try:
            import docx  # imported lazily: only .docx ingestion needs it
        except ImportError:
            return ParsedDocument(
                blocks=[],
                has_text_layer=None,
                parser=self.name,
                warnings=["python-docx is not installed; .docx files cannot be read."],
            )

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            return ParsedDocument(
                blocks=[],
                has_text_layer=None,
                parser=self.name,
                warnings=[f"DOCX could not be opened: {type(exc).__name__}"],
            )

        blocks: list[TextBlock] = []
        warnings: list[str] = []
        heading_stack: list[tuple[int, str]] = []
        offset = 0

        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style else ""
            level = _heading_level(style)
            if level:
                while heading_stack and heading_stack[-1][0] >= level:
                    heading_stack.pop()
                heading_stack.append((level, text))
                kind = "heading"
            else:
                kind = "prose"
            blocks.append(
                TextBlock(
                    text=text,
                    page=None,
                    char_start=offset,
                    char_end=offset + len(text),
                    section_path=" > ".join(h[1] for h in heading_stack) or None,
                    kind=kind,
                )
            )
            offset += len(text) + 1

        for table_index, table in enumerate(document.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header = rows[0]
            for row_index, row in enumerate(rows[1:], start=1):
                pairs = [f"{h}: {v}" for h, v in zip(header, row, strict=False) if v]
                if not pairs:
                    continue
                text = " | ".join(pairs)
                blocks.append(
                    TextBlock(
                        text=text,
                        page=None,
                        char_start=offset,
                        char_end=offset + len(text),
                        section_path=f"table {table_index + 1}",
                        kind="table_row",
                        metadata={"row_index": row_index, "columns": header},
                    )
                )
                offset += len(text) + 1

        if not blocks:
            warnings.append("Document contained no extractable text.")

        return ParsedDocument(
            blocks=blocks,
            page_count=None,
            has_text_layer=bool(blocks),
            parser=self.name,
            warnings=warnings,
            metadata={"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
        )


def _heading_level(style_name: str) -> int:
    if style_name.startswith("heading"):
        tail = style_name.replace("heading", "").strip()
        return int(tail) if tail.isdigit() else 1
    if style_name in {"title", "subtitle"}:
        return 1
    return 0
